#!/usr/bin/env python3
"""Génère js/17-simulateur-modele.js (formules du classeur exécutées par le site) à partir du
classeur Excel du simulateur.

Le site ne réécrit pas les calculs : il exécute les formules du classeur lui-même
(moteur js/16-simulateur-moteur.js). Ce script extrait toutes les formules et les
valeurs fixes du classeur dans un format compact :

  - chaque référence de cellule est réécrite en coordonnées absolues ($) ou relatives
    à la cellule qui la contient, si bien qu'une formule recopiée sur 25 colonnes
    (une par année) n'est stockée qu'une fois ;
  - format d'une référence dans un modèle : [feuille|ligne|colonne(|ligne2|colonne2)],
    feuille vide = feuille courante, nombre = absolu, ~n = décalage relatif.

Usage :  python3 tools/simulateur/build_model.py "chemin/Simulateur de Rentabilité_V4.xlsm"
Puis vérifier avec :  python3 tools/simulateur/check_model.py (comparaison au moteur de référence).
"""
import json, os, re, sys
import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from xl_oracle import tokenize, parse_ref, load_workbook  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', '..', 'js', '17-simulateur-modele.js')


def convert(formula, sheet, row, col, sheet_index):
    out = []
    for kind, text in tokenize(formula):
        if kind != 'ref':
            out.append(text)
            continue
        sh, r1, c1, r2, c2 = parse_ref(text, sheet)
        a1 = text.rsplit('!', 1)[-1].split(':')
        parts = []
        for (r, c), addr in zip(((r1, c1), (r2, c2)), a1 + a1[:1] if len(a1) == 1 else a1):
            m = re.match(r'(\$?)[A-Z]+(\$?)\d+', addr)
            col_abs, row_abs = m.group(1) == '$', m.group(2) == '$'
            parts.append(str(r) if row_abs else '~%d' % (r - row))
            parts.append(str(c) if col_abs else '~%d' % (c - col))
        if len(a1) == 1:
            parts = parts[:2]
        si = '' if sh == sheet else str(sheet_index[sh])
        out.append('[' + si + '|' + '|'.join(parts) + ']')
    return ''.join(out)


def build(path):
    wb = load_workbook(path)
    sheets = wb.sheetnames
    sheet_index = {s: i for i, s in enumerate(sheets)}
    templates, tpl_id = [], {}
    cells = []   # par feuille : [ligne, colonne, modèle, (1 si formule matricielle)]
    consts = []  # par feuille : [ligne, colonne, valeur]
    for ws in wb.worksheets:
        fc, kc = [], []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                arr = 0
                if hasattr(v, 'text'):
                    v, arr = v.text, 1
                if isinstance(v, str) and v.startswith('='):
                    t = convert(v[1:], ws.title, c.row, c.column, sheet_index)
                    if t not in tpl_id:
                        tpl_id[t] = len(templates)
                        templates.append(t)
                    fc.append([c.row, c.column, tpl_id[t]] + ([1] if arr else []))
                elif v is not None:
                    if isinstance(v, bool):
                        kc.append([c.row, c.column, v])
                    elif isinstance(v, (int, float)):
                        kc.append([c.row, c.column, float(v) if isinstance(v, float) else v])
                    elif isinstance(v, str):
                        kc.append([c.row, c.column, v])
        cells.append(fc)
        consts.append(kc)
    # Mise en forme des lignes (pour afficher les tableaux comme dans le classeur) :
    # [ligne, repliée (groupe masqué dans Excel), format des montants de la ligne]
    meta = []
    for ws in wb.worksheets:
        rows = []
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 2).value is None:
                continue
            d = ws.row_dimensions.get(r)
            hidden = 1 if (d is not None and (d.hidden or d.outlineLevel)) else 0
            rows.append([r, hidden, number_format_code(ws.cell(r, 4).number_format)])
        meta.append(rows)
    return {'sheets': sheets, 't': templates, 'f': cells, 'k': consts, 'm': meta}


def number_format_code(nf):
    """Réduit un format de nombre Excel à un code d'affichage : e (euros), p0/p2 (pourcentage),
    n1 (une décimale), i (entier), g (général)."""
    nf = nf or 'General'
    if '%' in nf:
        m = re.search(r'0\.(0+)%', nf)
        return 'p%d' % (len(m.group(1)) if m else 0)
    if re.search(r'#,##0\.0[^0]', nf + ' '):
        return 'n1'
    if '€' in nf or '_(*' in nf:
        return 'e'
    if nf.startswith('#,##0') or nf == '0':
        return 'i'
    return 'g'


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser('~/Desktop/Simulateur de Rentabilité_V4.xlsm')
    model = build(src)
    body = json.dumps(model, ensure_ascii=False, separators=(',', ':'))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        f.write('// FICHIER GÉNÉRÉ par tools/simulateur/build_model.py depuis « %s » — ne pas modifier à la main.\n' % os.path.basename(src))
        f.write('const SIM_MODEL = ' + body + ';\n')
    n = sum(len(x) for x in model['f'])
    print('%d formules, %d modèles distincts, %d octets -> %s' % (n, len(model['t']), len(body.encode()), os.path.relpath(OUT)))


if __name__ == '__main__':
    main()
