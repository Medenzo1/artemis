# Mini Excel formula evaluator for the "Simulateur de Rentabilité V4" workbook.
# Supports exactly the functions used by that workbook, with array semantics for array formulas.
import re, math, sys, threading
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

sys.setrecursionlimit(1000000)


def load_workbook(path, **kw):
    """openpyxl efface le contenu des cellules fusionnées (sauf celle en haut à gauche) ;
    Excel, lui, les calcule toujours. On conserve donc leur contenu."""
    from openpyxl.worksheet.worksheet import Worksheet
    orig = Worksheet._clean_merge_range
    Worksheet._clean_merge_range = lambda self, mcr: None
    try:
        return openpyxl.load_workbook(path, **kw)
    finally:
        Worksheet._clean_merge_range = orig
threading.stack_size(1024 * 1024 * 1024)


class XLErr(Exception):
    def __init__(self, code): self.code = code
    def __repr__(self): return self.code
    __str__ = __repr__
    def __eq__(self, o): return isinstance(o, XLErr) and o.code == self.code
    def __hash__(self): return hash(self.code)

NA = XLErr('#N/A'); VALUE = XLErr('#VALUE!'); DIV0 = XLErr('#DIV/0!'); REF = XLErr('#REF!'); NUM = XLErr('#NUM!'); NAME = XLErr('#NAME?')
ERRS = {e.code: e for e in (NA, VALUE, DIV0, REF, NUM, NAME)}


class Blank:
    def __repr__(self): return 'BLANK'
BLANK = Blank()


class Rng:
    __slots__ = ('sheet', 'r1', 'c1', 'r2', 'c2')
    def __init__(self, sheet, r1, c1, r2, c2): self.sheet, self.r1, self.c1, self.r2, self.c2 = sheet, r1, c1, r2, c2


class Arr:
    __slots__ = ('rows',)
    def __init__(self, rows): self.rows = rows  # list of lists


# ─────────────── tokenizer ───────────────
TOK = re.compile(r"""
 (?P<ws>\s+)
|(?P<str>"(?:[^"]|"")*")
|(?P<err>\#(?:N/A|VALUE!|DIV/0!|REF!|NUM!|NAME\?|NULL!))
|(?P<ref>(?:'(?:[^']|'')+'|[A-Za-z_][\w\.]*)?!?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
|(?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?|\.\d+)
|(?P<func>[A-Z][A-Z0-9\.]*\()
|(?P<bool>TRUE|FALSE)
|(?P<op><>|<=|>=|[-+*/^&=<>%(),:!])
""", re.X)

def tokenize(s):
    out = []; i = 0
    while i < len(s):
        m = TOK.match(s, i)
        if not m: raise SyntaxError('tok at %d: %r' % (i, s[i:i+30]))
        i = m.end(); k = m.lastgroup; v = m.group()
        if k == 'ws': continue
        if k == 'ref' and '!' not in v and not re.match(r"^\$?[A-Z]{1,3}\$?\d+(:\$?[A-Z]{1,3}\$?\d+)?$", v):
            raise SyntaxError('bad ref ' + v)
        out.append((k, v))
    return out


def parse_ref(v, cursheet):
    if '!' in v:
        sh, a = v.rsplit('!', 1)
        if sh.startswith("'"): sh = sh[1:-1].replace("''", "'")
    else:
        sh, a = cursheet, v
    parts = a.replace('$', '').split(':')
    def rc(p):
        m = re.match(r'([A-Z]+)(\d+)', p); return int(m.group(2)), column_index_from_string(m.group(1))
    r1, c1 = rc(parts[0])
    r2, c2 = rc(parts[1]) if len(parts) > 1 else (r1, c1)
    return sh, r1, c1, r2, c2


# ─────────────── parser (Pratt) → AST tuples ───────────────
BIN = {'=': 1, '<>': 1, '<': 1, '>': 1, '<=': 1, '>=': 1, '&': 2, '+': 3, '-': 3, '*': 4, '/': 4, '^': 5}

class Parser:
    def __init__(self, toks, sheet): self.t = toks; self.i = 0; self.sheet = sheet
    def peek(self): return self.t[self.i] if self.i < len(self.t) else (None, None)
    def nxt(self): x = self.peek(); self.i += 1; return x
    def expr(self, minp=0):
        lhs = self.unary()
        while True:
            k, v = self.peek()
            if k == 'op' and v == '%':
                self.nxt(); lhs = ('pct', lhs); continue
            if k == 'op' and v in BIN and BIN[v] >= minp and BIN[v] > 0:
                p = BIN[v]; self.nxt()
                rhs = self.expr(p + 1)  # left-assoc
                lhs = ('bin', v, lhs, rhs); continue
            return lhs
    def unary(self):
        k, v = self.peek()
        if k == 'op' and v in '+-':
            self.nxt(); e = self.unary()
            return ('neg', e) if v == '-' else ('pos', e)
        return self.primary()
    def primary(self):
        k, v = self.nxt()
        if k == 'num': return ('num', float(v))
        if k == 'str': return ('str', v[1:-1].replace('""', '"'))
        if k == 'bool': return ('bool', v == 'TRUE')
        if k == 'err': return ('err', ERRS.get(v, REF))
        if k == 'ref': return ('ref',) + parse_ref(v, self.sheet)
        if k == 'func':
            name = v[:-1]; args = []
            if self.peek() == ('op', ')'): self.nxt(); return ('fn', name, args)
            while True:
                if self.peek() in (('op', ','), ('op', ')')): args.append(('blankarg',))
                else: args.append(self.expr())
                k2, v2 = self.nxt()
                if v2 == ')': break
                if v2 != ',': raise SyntaxError('expected , got ' + str(v2))
            return ('fn', name, args)
        if k == 'op' and v == '(':
            e = self.expr(); k2, v2 = self.nxt()
            if v2 != ')': raise SyntaxError('expected )')
            return e
        raise SyntaxError('unexpected %r %r' % (k, v))


# ─────────────── value helpers ───────────────
NUMSTR = re.compile(r'^\s*[-+]?(\d[\d\s  ]*)([.,]\d+)?\s*(€|%)?\s*$')

def to_num(v):
    if isinstance(v, XLErr): raise v
    if v is BLANK or v is None: return 0.0
    if isinstance(v, bool): return 1.0 if v else 0.0
    if isinstance(v, (int, float)): return float(v)
    if isinstance(v, str):
        s = v.strip()
        try: return float(s)
        except ValueError: pass
        m = NUMSTR.match(s)
        if m:
            n = float(re.sub(r'[\s  ]', '', m.group(1)) + (m.group(2) or '').replace(',', '.'))
            if s.lstrip().startswith('-'): n = -n
            if m.group(3) == '%': n /= 100
            return n
        raise VALUE
    raise VALUE

def to_str(v):
    if isinstance(v, XLErr): raise v
    if v is BLANK or v is None: return ''
    if isinstance(v, bool): return 'TRUE' if v else 'FALSE'
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15: return str(int(v))
        return ('%.15g' % v).replace('.', ',')
    return str(v)

def to_bool(v):
    if isinstance(v, XLErr): raise v
    if v is BLANK or v is None: return False
    if isinstance(v, bool): return v
    if isinstance(v, (int, float)): return v != 0
    if isinstance(v, str):
        if v.upper() == 'TRUE': return True
        if v.upper() == 'FALSE': return False
        raise VALUE
    raise VALUE

def typerank(v):
    if isinstance(v, bool): return 2
    if isinstance(v, str): return 1
    return 0

def compare(a, b, op):
    if isinstance(a, XLErr): return a
    if isinstance(b, XLErr): return b
    if a is BLANK: a = '' if isinstance(b, str) else (False if isinstance(b, bool) else 0.0)
    if b is BLANK: b = '' if isinstance(a, str) else (False if isinstance(a, bool) else 0.0)
    ra, rb = typerank(a), typerank(b)
    if ra != rb: c = (ra > rb) - (ra < rb)
    elif ra == 1: x, y = a.lower(), b.lower(); c = (x > y) - (x < y)
    else: x, y = float(a), float(b); c = (x > y) - (x < y)
    return {'=': c == 0, '<>': c != 0, '<': c < 0, '>': c > 0, '<=': c <= 0, '>=': c >= 0}[op]

def xround(x, n, mode):
    d = Decimal(repr(float(x)))
    q = Decimal(1).scaleb(-int(n))
    return float(d.quantize(q, rounding=mode))


# ─────────────── workbook model ───────────────
class Book:
    def __init__(self, path):
        wb = load_workbook(path)
        wv = load_workbook(path, data_only=True)
        self.sheets = wb.sheetnames
        self.formula = {}   # (sheet,r,c) -> (ast, is_array)
        self.const = {}     # (sheet,r,c) -> value
        self.cached = {}
        for ws in wb.worksheets:
            vs = wv[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    key = (ws.title, c.row, c.column)
                    cv = vs.cell(c.row, c.column).value
                    if cv is not None: self.cached[key] = cv
                    isarr = False
                    if hasattr(v, 'text'): v = v.text; isarr = True
                    if isinstance(v, str) and v.startswith('='):
                        self.formula[key] = (v, isarr)
                    elif v is not None:
                        self.const[key] = float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else v
        self.ast = {}
        for key, (f, isarr) in self.formula.items():
            self.ast[key] = (Parser(tokenize(f[1:]), key[0]).expr(), isarr)
        self.over = {}
        self.reset()

    def reset(self):
        self.memo = {}

    def set(self, sheet, addr, value):
        m = re.match(r'([A-Z]+)(\d+)', addr)
        key = (sheet, int(m.group(2)), column_index_from_string(m.group(1)))
        self.over[key] = BLANK if value is None else (float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else value)
        self.memo = {}

    def get(self, sheet, addr):
        m = re.match(r'([A-Z]+)(\d+)', addr)
        return self.cell((sheet, int(m.group(2)), column_index_from_string(m.group(1))))

    def cell(self, key):
        if key in self.over: return self.over[key]
        if key in self.memo:
            v = self.memo[key]
            if v is _INPROG: raise RuntimeError('circular at %r' % (key,))
            return v
        if key in self.ast:
            self.memo[key] = _INPROG
            ast, isarr = self.ast[key]
            try:
                v = self.ev(ast, key, isarr)
                if isinstance(v, Rng): v = self.implicit(v, key)
                if isinstance(v, Arr): v = v.rows[0][0]
                if v is BLANK: v = 0.0
            except XLErr as e:
                v = e
            self.memo[key] = v
            return v
        return self.const.get(key, BLANK)

    def implicit(self, r, at):
        if r.r1 == r.r2 and r.c1 == r.c2: return self.cell((r.sheet, r.r1, r.c1))
        if r.r1 == r.r2 and r.c1 <= at[2] <= r.c2: return self.cell((r.sheet, r.r1, at[2]))
        if r.c1 == r.c2 and r.r1 <= at[1] <= r.r2: return self.cell((r.sheet, at[1], r.c1))
        raise VALUE

    def rng_rows(self, r):
        return [[self.cell((r.sheet, i, j)) for j in range(r.c1, r.c2 + 1)] for i in range(r.r1, r.r2 + 1)]

    def grid(self, v):
        if isinstance(v, Rng): return self.rng_rows(v)
        if isinstance(v, Arr): return v.rows
        return [[v]]

    def flat(self, v):
        return [x for row in self.grid(v) for x in row]

    # scalarize value in non-array context
    def scal(self, v, at, isarr):
        if isinstance(v, Rng):
            if isarr: return Arr(self.rng_rows(v))
            return self.implicit(v, at)
        return v

    def ev(self, n, at, isarr):
        t = n[0]
        if t == 'num' or t == 'str' or t == 'bool': return n[1]
        if t == 'err': raise n[1]
        if t == 'blankarg': return BLANK
        if t == 'ref':
            _, sh, r1, c1, r2, c2 = n
            if r1 == r2 and c1 == c2: return self.cell((sh, r1, c1))
            return Rng(sh, r1, c1, r2, c2)
        if t == 'neg': return self.unop(lambda x: -to_num(x), self.scal(self.ev(n[1], at, isarr), at, isarr))
        if t == 'pos': return self.scal(self.ev(n[1], at, isarr), at, isarr)
        if t == 'pct': return self.unop(lambda x: to_num(x) / 100, self.scal(self.ev(n[1], at, isarr), at, isarr))
        if t == 'bin':
            op = n[1]
            a = self.scal(self.ev(n[2], at, isarr), at, isarr)
            b = self.scal(self.ev(n[3], at, isarr), at, isarr)
            return self.binop(op, a, b)
        if t == 'fn': return self.call(n[1], n[2], at, isarr)
        raise RuntimeError(t)

    def unop(self, f, a):
        if isinstance(a, Arr): return Arr([[self._safe(f, x) for x in row] for row in a.rows])
        return f(a)

    def _safe(self, f, *xs):
        try: return f(*xs)
        except XLErr as e: return e

    def binop(self, op, a, b):
        if isinstance(a, Arr) or isinstance(b, Arr):
            ga = a.rows if isinstance(a, Arr) else [[a]]
            gb = b.rows if isinstance(b, Arr) else [[b]]
            R = max(len(ga), len(gb)); C = max(len(ga[0]), len(gb[0]))
            def at(g, i, j):
                ii = 0 if len(g) == 1 else i; jj = 0 if len(g[0]) == 1 else j
                if ii >= len(g) or jj >= len(g[0]): return NA
                return g[ii][jj]
            return Arr([[self._safe(lambda x, y: self.binop(op, x, y), at(ga, i, j), at(gb, i, j)) for j in range(C)] for i in range(R)])
        if isinstance(a, XLErr): raise a
        if isinstance(b, XLErr): raise b
        if op in ('=', '<>', '<', '>', '<=', '>='): return compare(a, b, op)
        if op == '&': return to_str(a) + to_str(b)
        x, y = to_num(a), to_num(b)
        if op == '+': return x + y
        if op == '-': return x - y
        if op == '*': return x * y
        if op == '/':
            if y == 0: raise DIV0
            return x / y
        if op == '^':
            try: return math.pow(x, y)
            except (ValueError, OverflowError): raise NUM

    # ─────────── functions ───────────
    def call(self, name, args, at, isarr):
        E = lambda a: (Rng(a[1], a[2], a[3], a[4], a[5]) if a[0] == 'ref' else self.ev(a, at, isarr))
        S = lambda a: self.scal(self.ev(a, at, isarr), at, isarr)
        if name == 'IF':
            c = S(args[0])
            if isinstance(c, Arr):
                t = S(args[1]) if len(args) > 1 else True
                f = S(args[2]) if len(args) > 2 else False
                gt = t.rows if isinstance(t, Arr) else None
                gf = f.rows if isinstance(f, Arr) else None
                rows = []
                for i, row in enumerate(c.rows):
                    out = []
                    for j, cv in enumerate(row):
                        try:
                            b = to_bool(cv)
                            if b: out.append(gt[i if len(gt) > 1 else 0][j if len(gt[0]) > 1 else 0] if gt else t)
                            else: out.append(gf[i if len(gf) > 1 else 0][j if len(gf[0]) > 1 else 0] if gf else f)
                        except XLErr as e: out.append(e)
                    rows.append(out)
                return Arr(rows)
            if to_bool(c):
                return S(args[1]) if len(args) > 1 and args[1][0] != 'blankarg' else (True if len(args) < 2 else 0.0)
            if len(args) > 2: return S(args[2]) if args[2][0] != 'blankarg' else 0.0
            return False
        if name == 'IFERROR':
            try:
                v = S(args[0])
                if isinstance(v, XLErr): raise v
                if isinstance(v, Arr):
                    alt = None
                    rows = []
                    for row in v.rows:
                        o = []
                        for x in row:
                            if isinstance(x, XLErr):
                                if alt is None: alt = S(args[1])
                                o.append(alt)
                            else: o.append(x)
                        rows.append(o)
                    return Arr(rows)
                return v
            except XLErr:
                return S(args[1])
        if name in ('AND', 'OR'):
            vals = []
            for a in args:
                v = E(a)
                if isinstance(v, (Rng, Arr)):
                    for x in self.flat(v):
                        if isinstance(x, XLErr): raise x
                        if isinstance(x, (bool, float, int)) and not isinstance(x, str): vals.append(to_bool(x))
                else: vals.append(to_bool(v))
            if not vals: raise VALUE
            return all(vals) if name == 'AND' else any(vals)
        if name in ('SUM', 'MAX', 'MIN'):
            nums = []
            for a in args:
                v = E(a)
                if isinstance(v, (Rng, Arr)):
                    for x in self.flat(v):
                        if isinstance(x, XLErr): raise x
                        if isinstance(x, (int, float)) and not isinstance(x, bool): nums.append(float(x))
                else:
                    nums.append(to_num(v))
            if name == 'SUM': return float(sum(nums))
            if not nums: return 0.0
            return max(nums) if name == 'MAX' else min(nums)
        if name == 'POWER':
            return self.binop('^', S(args[0]), S(args[1]))
        if name in ('ROUND', 'ROUNDDOWN'):
            x = to_num(S(args[0])); d = to_num(S(args[1])) if len(args) > 1 else 0
            return xround(x, d, ROUND_HALF_UP if name == 'ROUND' else ROUND_DOWN)
        if name in ('LEFT', 'RIGHT'):
            s = to_str(S(args[0])); k = int(to_num(S(args[1]))) if len(args) > 1 else 1
            return s[:k] if name == 'LEFT' else (s[-k:] if k else '')
        if name in ('HLOOKUP', 'VLOOKUP'):
            key = S(args[0]); tv = E(args[1]); idx = int(to_num(S(args[2])))
            approx = True
            if len(args) > 3:
                a3 = S(args[3]); approx = to_bool(a3) if a3 is not BLANK else False
            if isinstance(key, XLErr): raise key
            if isinstance(tv, Rng):
                r = tv
                if name == 'HLOOKUP':
                    if idx > r.r2 - r.r1 + 1: raise REF
                    keys = [self.cell((r.sheet, r.r1, j)) for j in range(r.c1, r.c2 + 1)]
                    j = self.lookup(key, keys, approx); return self.cell((r.sheet, r.r1 + idx - 1, r.c1 + j))
                else:
                    if idx > r.c2 - r.c1 + 1: raise REF
                    keys = [self.cell((r.sheet, i, r.c1)) for i in range(r.r1, r.r2 + 1)]
                    j = self.lookup(key, keys, approx); return self.cell((r.sheet, r.r1 + j, r.c1 + idx - 1))
            tbl = self.grid(tv)
            if name == 'HLOOKUP':
                keys = tbl[0]; j = self.lookup(key, keys, approx); return tbl[idx - 1][j]
            keys = [row[0] for row in tbl]; j = self.lookup(key, keys, approx); return tbl[j][idx - 1]
        if name == 'MATCH':
            key = S(args[0]); arr = self.flat(E(args[1]))
            mt = to_num(S(args[2])) if len(args) > 2 else 1
            if isinstance(key, XLErr): raise key
            if mt == 0: return float(self.lookup(key, arr, False) + 1)
            if mt == 1: return float(self.lookup(key, arr, True) + 1)
            raise RuntimeError('MATCH -1')
        if name == 'INDEX':
            tv = E(args[0])
            r = int(to_num(S(args[1]))) if len(args) > 1 and args[1][0] != 'blankarg' else 0
            c = int(to_num(S(args[2]))) if len(args) > 2 and args[2][0] != 'blankarg' else 0
            if isinstance(tv, Rng):
                R, C = tv.r2 - tv.r1 + 1, tv.c2 - tv.c1 + 1
            else:
                g = self.grid(tv); R, C = len(g), len(g[0])
            if R == 1 and len(args) == 2: c, r = r, 1
            if r == 0: r = 1
            if c == 0: c = 1
            if r > R or c > C or r < 1 or c < 1: raise REF
            v = self.cell((tv.sheet, tv.r1 + r - 1, tv.c1 + c - 1)) if isinstance(tv, Rng) else g[r - 1][c - 1]
            if isinstance(v, XLErr): raise v
            return v
        if name == 'SUMIFS':
            sr = self.flat(E(args[0]))
            mask = [True] * len(sr)
            for k in range(1, len(args), 2):
                cr = self.flat(E(args[k])); crit = S(args[k + 1])
                test = self.crit(crit)
                for i, x in enumerate(cr): mask[i] = mask[i] and test(x)
            return float(sum(float(x) for x, m in zip(sr, mask) if m and isinstance(x, (int, float)) and not isinstance(x, bool)))
        if name == 'SUMPRODUCT':
            grids = []
            for a in args:
                v = E(a)
                if isinstance(v, Rng): v = Arr(self.rng_rows(v))
                elif not isinstance(v, Arr): v = Arr([[v]])
                grids.append(v.rows)
            tot = 0.0
            R, C = len(grids[0]), len(grids[0][0])
            for g in grids:
                if len(g) != R or len(g[0]) != C: raise VALUE
            for i in range(R):
                for j in range(C):
                    p = 1.0
                    for g in grids:
                        x = g[i][j]
                        if isinstance(x, XLErr): raise x
                        p *= float(x) if isinstance(x, (int, float)) and not isinstance(x, bool) else 0.0
                    tot += p
            return tot
        if name == 'NPV':
            rate = to_num(S(args[0])); vals = []
            for a in args[1:]:
                v = E(a)
                if isinstance(v, (Rng, Arr)):
                    for x in self.flat(v):
                        if isinstance(x, XLErr): raise x
                        if isinstance(x, (int, float)) and not isinstance(x, bool): vals.append(float(x))
                else: vals.append(to_num(v))
            return sum(x / (1 + rate) ** (i + 1) for i, x in enumerate(vals))
        if name == 'IRR':
            vals = []
            for x in self.flat(E(args[0])):
                if isinstance(x, XLErr): raise x
                if isinstance(x, (int, float)) and not isinstance(x, bool): vals.append(float(x))
            guess = to_num(S(args[1])) if len(args) > 1 else 0.1
            return irr(vals, guess)
        raise RuntimeError('unsupported function ' + name)

    def crit(self, c):
        if isinstance(c, str):
            m = re.match(r'^(<=|>=|<>|<|>|=)?(.*)$', c)
            op = m.group(1) or '='; rhs = m.group(2)
            try: rv = float(rhs)
            except ValueError: rv = rhs
            if isinstance(rv, str):
                return lambda x: compare(x if isinstance(x, str) else to_str(x), rv, op) if not isinstance(x, XLErr) else False
            return lambda x: (isinstance(x, (int, float)) and not isinstance(x, bool) and compare(float(x), rv, op))
        if isinstance(c, XLErr): raise c
        return lambda x: (not isinstance(x, XLErr)) and x is not BLANK and compare(x, c, '=') is True

    def lookup(self, key, keys, approx):
        if not approx:
            for j, k in enumerate(keys):
                if isinstance(k, XLErr) or k is BLANK: continue
                if typerank(k) == typerank(key) and compare(k, key, '=') is True: return j
            raise NA
        # approximate: Excel binary search
        lo, hi = 0, len(keys) - 1; best = -1
        # emulate Excel binary search over same-typed values
        while lo <= hi:
            mid = (lo + hi) // 2
            k = keys[mid]
            # skip blanks/other types by moving
            if k is BLANK or isinstance(k, XLErr) or typerank(k) != typerank(key):
                # linear fallback in this rare case
                cand = [j for j, kk in enumerate(keys) if not (kk is BLANK or isinstance(kk, XLErr)) and typerank(kk) == typerank(key) and compare(kk, key, '<=') is True]
                if not cand: raise NA
                return cand[-1]
            if compare(k, key, '<=') is True: best = mid; lo = mid + 1
            else: hi = mid - 1
        if best < 0: raise NA
        return best


def irr(vals, guess=0.1, maxit=20, tol=1e-7):
    """TRI compatible Excel (algorithme reconstitué par comparaison avec Excel, l'original n'étant pas public) :
    1) Newton sur x = 1/(1+taux) depuis l'estimation ; 2) en cas d'échec, Newton sur le taux depuis -10 %.
    20 itérations au plus, précision 1e-7 ; un taux <= -100 % est rejeté (#NOMBRE!)."""
    def newton_x(g):
        x = 1 / (1 + g)
        for _ in range(maxit):
            f = 0.0; d = 0.0
            for i, v in enumerate(vals):
                f += v * x ** i
                if i: d += i * v * x ** (i - 1)
            if d == 0: return None
            nx = x - f / d
            if nx == 0 or not math.isfinite(nx): return None
            r0 = 1 / x - 1; r1 = 1 / nx - 1
            if abs(r1 - r0) <= tol: return r1 if r1 > -1 else None
            x = nx
        return None
    def newton_r(g):
        r = g
        for _ in range(maxit):
            try:
                f = 0.0; d = 0.0
                for i, v in enumerate(vals):
                    f += v / (1 + r) ** i
                    d -= i * v / (1 + r) ** (i + 1)
            except (ZeroDivisionError, OverflowError):
                return None
            if isinstance(f, complex) or d == 0 or not math.isfinite(f) or not math.isfinite(d): return None
            nr = r - f / d
            if not math.isfinite(nr): return None
            if abs(nr - r) <= tol: return nr if nr > -1 else None
            r = nr
        return None
    try:
        v = newton_x(guess)
    except (ZeroDivisionError, OverflowError):
        v = None
    if v is None: v = newton_r(-0.1)
    if v is None: raise NUM
    return v


class _InProg: pass
_INPROG = _InProg()


def run_threaded(fn):
    res = {}
    def w():
        try: res['v'] = fn()
        except BaseException as e: res['e'] = e
    t = threading.Thread(target=w); t.start(); t.join()
    if 'e' in res: raise res['e']
    return res['v']
